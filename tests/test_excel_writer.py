from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from copy import copy

from datetime import date, datetime
from decimal import Decimal

from tempfile import TemporaryDirectory

from unittest.mock import patch

from pathlib import Path

from axelpro_riport.excel_writer import (
    copy_block,
    clear_input_values,
    create_week_block,
    create_year_workbook,
    load_or_create_year_workbook,
    set_week_dates,
    create_month_sheet,
    ensure_month_exchange_rate,
    ensure_month_sheet,
    ensure_summary_month,
    ensure_months_through,
    find_day_column,
    add_cell_amount,
    write_daily_totals,
    write_monthly_totals,
    write_import_result,
    get_year_workbook_path,
    create_workbook_backup,
    save_workbook_safely,
    PROCESSED_DOCUMENTS_SHEET,
    ensure_processed_documents_sheet,
    read_processed_documents,
    add_processed_documents,
    DAILY_EXCHANGE_RATES_SHEET,
    add_daily_exchange_rates,
    read_daily_exchange_rates,
    update_monthly_exchange_rates,
)

from axelpro_riport.calendar_utils import get_month_weeks

from axelpro_riport.models import DailyTotals, MonthlyTotals, ImportResult, DailyExchangeRate


def test_copy_block() -> None:
    workbook = Workbook()
    source = workbook.active
    source.title = "Forrás"
    target = workbook.create_sheet(title="Cél")

    source["A1"] = "Teszt"
    source["B2"] = 100
    source["B3"] = "=B2+$B$2"

    source["A1"].font = Font(bold=True)
    source["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    source["B2"].number_format = "#,##0"

    source.row_dimensions[1].height = 25
    source.column_dimensions["A"].width = 20
    source.merge_cells("A1:C1")

    copy_block(
        source_sheet=source,
        target_sheet=target,
        min_row=1,
        max_row=3,
        min_column=1,
        max_column=3,
        target_start_row=5,
        target_start_column=5
    )

    assert target["E5"].value == "Teszt"
    assert target["F6"].value == 100
    assert target["F7"].value == "=F6+$B$2"

    assert "E5:G5" in {str(merged_range) for merged_range in target.merged_cells.ranges}

    assert copy(target["E5"].font) == copy(source["A1"].font)
    assert copy(target["E5"].fill) == copy(source["A1"].fill)
    assert target["F6"].number_format == "#,##0"

    assert target.row_dimensions[5].height == 25
    assert target.column_dimensions["E"].width == 20

    workbook.close()
    print("Sikeres blokkmásolási teszt.")

def test_clear_input_values() -> None:
    workbook = Workbook()
    sheet = workbook.active

    sheet["A1"] = "Sorfelirat"
    sheet["B1"] = 150
    sheet.merge_cells("B1:C1")

    sheet["B2"] = "=B1*2"
    sheet["C2"] = 200
    sheet["B1"].font = Font(bold=True)

    clear_input_values(
        sheet=sheet,
        min_row=1,
        max_row=2,
        min_column=2,
        max_column=3
    )

    assert sheet["A1"].value == "Sorfelirat"
    assert sheet["B1"].value is None
    assert sheet["C2"].value is None
    assert sheet["B1"].font.bold is True
    assert "B1:C1" in {str(merged_range) for merged_range in sheet.merged_cells.ranges}

    workbook.close()
    print("Sikeres adatürítési teszt.")

def test_create_week_block() -> None:
    workbook = Workbook()
    template = workbook.active
    target = workbook.create_sheet(title="Cél")

    template["A8"] = "Bontott alkatrész sz."

    day_names = (
        "Hétfő",
        "Kedd",
        "Szerda",
        "Csütörtök",
        "Péntek",
        "Szombat"
    )

    for weekday, day_name in enumerate(day_names):
        column = 2 + weekday * 2

        template.cell(row=1, column=column).value = day_name
        template.cell(row=2, column=column).value = date(2026, 1, 1)
        template.cell(row=4, column=column).value = 100
        template.cell(row=8, column=column).value = 200

        for row in (1, 2, 4, 8):
            template.merge_cells(
                start_row=row,
                end_row=row,
                start_column=column,
                end_column=column + 1
            )

    template["N1"] = "Heti záró"
    template["N8"] = "=B8+D8+F8+H8+J8+L8"
    template["N26"] = "=B26+D26+F26+H26+J26+L26"
    template["O26"] = "=C26+E26+G26+I26+K26+M26"

    for row in (1, 2, 4, 8):
        template.merge_cells(
            start_row=row,
            end_row=row,
            start_column=14,
            end_column=15
        )

    first_width = create_week_block(
        template_sheet=template,
        target_sheet=target,
        week=get_month_weeks(2026, 10)[0],
        week_start_column=1,
        include_labels=True
    )

    assert first_width == 9
    assert target["A8"].value == "Bontott alkatrész sz."
    assert target["B1"].value == "Csütörtök"
    assert target["D1"].value == "Péntek"
    assert target["F1"].value == "Szombat"
    assert target["H1"].value == "Heti záró"

    assert target["B2"].value == date(2026, 10, 1)
    assert target["D2"].value == date(2026, 10, 2)
    assert target["F2"].value == date(2026, 10, 3)
    assert target["H2"].value == 40
    assert target["B4"].value is None
    assert target["B8"].value is None

    assert target["H8"].value == "=B8+D8+F8"
    assert target["H26"].value == "=B26+D26+F26"
    assert target["I26"].value == "=C26+E26+G26"

    next_column = 1 + first_width

    last_width = create_week_block(
        template_sheet=template,
        target_sheet=target,
        week=get_month_weeks(2026, 9)[-1],
        week_start_column=next_column,
        include_labels=False
    )

    assert last_width == 8
    assert target["J1"].value == "Hétfő"
    assert target["L1"].value == "Kedd"
    assert target["N1"].value == "Szerda"
    assert target["P1"].value == "Heti záró"

    assert target["J2"].value == date(2026, 9, 28)
    assert target["L2"].value == date(2026, 9, 29)
    assert target["N2"].value == date(2026, 9, 30)
    assert target["P2"].value == 40

    assert target["P8"].value == "=J8+L8+N8"
    assert target["P26"].value == "=J26+L26+N26"
    assert target["Q26"].value == "=K26+M26+O26"

    merged_ranges = {str(merged_range) for merged_range in target.merged_cells.ranges}

    assert "B8:C8" in merged_ranges
    assert "H8:I8" in merged_ranges
    assert "J8:K8" in merged_ranges
    assert "P8:Q8" in merged_ranges

    next_column = next_column + last_width

    sunday_width = create_week_block(
        template_sheet=template,
        target_sheet=target,
        week=[date(2026, 11, 1)],
        week_start_column=next_column,
        include_labels=False
    )

    assert sunday_width == 0
    assert target.max_column == next_column - 1

    workbook.close()
    print("Sikeres rövidített heti blokk teszt.")

def test_set_week_dates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    week = get_month_weeks(2026, 10)[0]

    set_week_dates(
        sheet=sheet,
        week=week,
        week_start_column=1,
        include_labels=True
    )

    assert sheet["B2"].value == date(2026, 10, 1)
    assert sheet["D2"].value == date(2026, 10, 2)
    assert sheet["F2"].value == date(2026, 10, 3)
    assert sheet["H2"].value == 40

    set_week_dates(
        sheet=sheet,
        week=week,
        week_start_column=10,
        include_labels=False
    )

    assert sheet["J2"].value == date(2026, 10, 1)
    assert sheet["L2"].value == date(2026, 10, 2)
    assert sheet["N2"].value == date(2026, 10, 3)
    assert sheet["P2"].value == 40

    workbook.close()
    print("Sikeres heti dátumbeállítási teszt.")

def test_create_month_sheet() -> None:
    template_workbook = Workbook()
    template = template_workbook.active

    day_names = (
        "Hétfő",
        "Kedd",
        "Szerda",
        "Csütörtök",
        "Péntek",
        "Szombat"
    )

    template["A8"] = "Bontott alkatrész sz."

    for weekday, day_name in enumerate(day_names):
        column = 2 + weekday * 2
        template.cell(row=1, column=column).value = day_name

    template["N1"] = "Heti záró"
    template["N12"] = "=B12+D12+F12+H12+J12+L12"
    template["N26"] = "=B26+D26+F26+H26+J26+L26"
    template["O26"] = "=C26+E26+G26+I26+K26+M26"

    template["A43"] = "Dátum:"
    template["B43"] = "Havi záró"
    template["B44"] = "Szeptember"
    template["B45"] = "Forint"
    template["B53"] = "=B52+B48"
    template["B74"] = 123
    template["B78"] = "=B74*0.08"

    for row in range(43, 86):
        template.merge_cells(
            start_row=row,
            end_row=row,
            start_column=2,
            end_column=3
        )

    workbook = Workbook()
    workbook.remove(workbook.active)

    rate_reference = "'Euro'!$B$11"

    sheet = create_month_sheet(
        workbook=workbook,
        template_sheet=template,
        year=2026,
        month=10,
        exchange_rate_reference=rate_reference
    )

    assert sheet.title == "Október"
    assert sheet.max_column == 65

    assert sheet["B1"].value == "Csütörtök"
    assert sheet["B2"].value == date(2026, 10, 1)
    assert sheet["H2"].value == 40

    assert sheet["J1"].value == "Hétfő"
    assert sheet["J2"].value == date(2026, 10, 5)
    assert sheet["BJ2"].value == date(2026, 10, 31)
    assert sheet["BL2"].value == 44

    assert sheet["B43"].value == "Havi záró"
    assert sheet["B44"].value == "Október"
    assert sheet["B45"].value == "Forint"

    assert sheet["B46"].value == "=(H12+V12+AJ12+AX12+BL12)"
    assert sheet["B58"].value == "=(H26+V26+AJ26+AX26+BL26)+((I26+W26+AK26+AY26+BM26)*'Euro'!$B$11)"

    assert sheet["B59"].value == "=(H27+V27+AJ27+AX27+BL27)-(B61+B62+B63+B66+B67+B68)"

    assert sheet["B53"].value == "=B52+B48"
    assert sheet["B74"].value is None
    assert sheet["B78"].value == "=B74*0.08"

    assert "B46:C46" in {str(merged_range) for merged_range in sheet.merged_cells.ranges}

    try:
        create_month_sheet(
            workbook=workbook,
            template_sheet=template,
            year=2026,
            month=10,
            exchange_rate_reference=rate_reference
        )
    except ValueError:
        pass
    else:
        raise AssertionError(
            "A már létező havi munkalap létrehozása nem jelzett hibát."
        )

    workbook.close()
    template_workbook.close()
    print("Sikeres teljes havi munkalap létrehozási teszt.")

def test_ensure_month_sheet() -> None:
    template_workbook = Workbook()
    template = template_workbook.active

    template["N26"] = "=B26+D26+F26+H26+J26+L26"
    template["O26"] = "=C26+E26+G26+I26+K26+M26"

    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet = ensure_month_sheet(
        workbook=workbook,
        template_sheet=template,
        year=2026,
        month=10,
    )

    assert set(workbook.sheetnames) == {"Euro", "Október"}
    assert workbook["Euro"]["A11"].value == "Október"
    assert workbook["Euro"]["B11"].value is None
    assert workbook["Euro"]["B11"].number_format == "0.00"

    assert sheet["B58"].value == (
        "=(H26+V26+AJ26+AX26+BL26)"
        "+((I26+W26+AK26+AY26+BM26)*'Euro'!$B$11)"
    )

    sheet["B8"] = 12345
    workbook["Euro"]["B11"] = 400.25

    repeated_sheet = ensure_month_sheet(
        workbook=workbook,
        template_sheet=template,
        year=2026,
        month=10,
    )

    assert repeated_sheet is sheet
    assert repeated_sheet["B8"].value == 12345
    assert workbook["Euro"]["B11"].value == 400.25
    assert len(workbook.sheetnames) == 2

    reference = ensure_month_exchange_rate(
        workbook=workbook,
        month=10,
    )

    assert reference == "'Euro'!$B$11"
    assert workbook["Euro"]["B11"].value == 400.25

    workbook.close()
    template_workbook.close()
    print("Sikeres havi munkalap- és árfolyammegőrzési teszt.")

def test_summary_month() -> None:
    template_workbook = Workbook()
    template = template_workbook.active
    template["A46"] = "Alkatrész értékesítés bontott sz.n.:"
    template["A74"] = "Havi óraszám:"

    workbook = Workbook()
    workbook.remove(workbook.active)

    month_sheet = ensure_month_sheet(
        workbook=workbook,
        template_sheet=template,
        year=2026,
        month=1,
    )

    ensure_summary_month(
        workbook=workbook,
        template_sheet=template,
        month=1,
    )

    summary = workbook["Összesítő"]

    assert workbook.sheetnames[0] == "Összesítő"
    assert summary["B2"].value == "Január"
    assert summary["A4"].value == template["A46"].value

    assert summary["B4"].value == (
        '=IF(\'Január\'!B46="","",\'Január\'!B46)'
    )
    assert summary["B32"].value == (
        '=IF(\'Január\'!B74="","",\'Január\'!B74)'
    )

    assert summary["B7"].value is None
    assert summary.column_dimensions["B"].width == 22

    month_sheet["B74"] = 150

    ensure_summary_month(
        workbook=workbook,
        template_sheet=template,
        month=1,
    )

    assert workbook["Összesítő"] is summary
    assert month_sheet["B74"].value == 150
    assert workbook.sheetnames.count("Összesítő") == 1

    try:
        ensure_summary_month(
            workbook=workbook,
            template_sheet=template,
            month=2,
        )
    except ValueError:
        pass
    else:
        raise AssertionError(
            "Hiányzó havi munkalap esetén nem jelzett hibát."
        )

    assert summary.max_column == 2

    workbook.close()
    template_workbook.close()
    print("Sikeres összesítőmunkalap-teszt.")

def test_year_workbook() -> None:
    template_workbook = Workbook()
    template = template_workbook.active

    workbook = create_year_workbook(
        template_sheet=template,
        year=2026,
        last_month=2,
    )

    assert workbook.sheetnames == [
        "Összesítő",
        "Euro",
        "Január",
        "Február",
    ]

    assert workbook["Összesítő"]["B2"].value == "Január"
    assert workbook["Összesítő"]["C2"].value == "Február"
    assert workbook["Összesítő"].max_column == 3

    assert workbook["Euro"]["A2"].value == "Január"
    assert workbook["Euro"]["A3"].value == "Február"

    january = workbook["Január"]
    january["B8"] = 12345
    workbook["Euro"]["B2"] = 400.25

    ensure_months_through(
        workbook=workbook,
        template_sheet=template,
        year=2026,
        last_month=3,
    )

    assert workbook.sheetnames == [
        "Összesítő",
        "Euro",
        "Január",
        "Február",
        "Március",
    ]

    assert workbook["Január"] is january
    assert january["B8"].value == 12345
    assert workbook["Euro"]["B2"].value == 400.25

    assert workbook["Összesítő"]["D2"].value == "Március"
    assert workbook["Euro"]["A4"].value == "Március"

    new_year = create_year_workbook(
        template_sheet=template,
        year=2027,
        last_month=1,
    )

    assert new_year.sheetnames == [
        "Összesítő",
        "Euro",
        "Január",
    ]

    assert new_year["Összesítő"].max_column == 2
    assert new_year["Január"]["B8"].value is None
    assert new_year["Euro"]["B2"].value is None

    workbook.close()
    new_year.close()
    template_workbook.close()
    print("Sikeres éves munkafüzet- és hónapbővítési teszt.")

def test_find_day_column() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Október"

    sheet["B2"] = date(2026, 10, 1)
    sheet["D2"] = datetime(2026, 10, 2)
    sheet["H2"] = 40 # Heti sorszám, nem dátum.

    assert find_day_column(sheet, date(2026, 10, 1)) == 2
    assert find_day_column(sheet, date(2026, 10, 2)) == 4

    try:
        find_day_column(sheet, date(2026, 10, 3))
    except ValueError as error:
        assert "2026-10-03" in str(error)
    else:
        raise AssertionError("Hiányzó dátumnál hibát kellett volna jeleznie.")

    workbook.close()
    print("A dátumoszlop keresésének tesztje sikeres.")

def test_add_cell_amount() -> None:
    workbook = Workbook()
    sheet = workbook.active

    # Üres, összevont cálcellába írunk.
    sheet.merge_cells("B8:C8")
    add_cell_amount(sheet, row=8, column=2, amount=1200)
    assert sheet["B8"].value == 1200

    # Az új import növeli a meglévő értéket.
    add_cell_amount(sheet, row=8, column=2, amount=3500)
    assert sheet["B8"].value == 4700

    # A negatív korrekció csökkenti az összeget.
    add_cell_amount(sheet, row=8, column=2, amount=-700)
    assert sheet["B8"].value == 4000

    # Hibás tartalom esetén az eredeti érték megmarad.
    for invalid_value in ("szöveg", "=SUM(B8:C8)", True):
        sheet["B9"] = invalid_value

        try:
            add_cell_amount(sheet, row=9, column=2, amount=100)
        except ValueError as error:
            assert sheet["B9"].value == invalid_value
        else:
            raise AssertionError(
                "A hibás cellatartalom nem okozott hibát."
            )

    workbook.close()
    print("Az összegek hozzáadásának tesztje sikeres.")

def test_write_daily_totals() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["D2"] = date(2026, 10, 2)

    expected = {
        8: 14000,
        9: 37000,
        10: 22500,
        23: 10000,
        24: 33000,
        27: 13125,
        32: 1250,
        40: 7500,
    }

    for row in expected:
        sheet.merge_cells(
            start_row=row,
            end_row=row,
            start_column=4,
            end_column=5
        )

    sheet["D8"] = 1000
    sheet["D4"] = 900 # Kézzel kitöltött, nem célcella.
    sheet["B8"] = 500 # Másik napi adat.
    sheet["B14"] = "=B12+B13"

    totals = DailyTotals(
        movement_date=date(2026, 10, 2),
        used_parts=14000,
        new_parts=37000,
        labor=22500,
        workshop_used_parts=10000,
        workshop_new_parts=33000,
        purchases=13125,
        shipping=1250,
        trailer=7500,
    )

    write_daily_totals(sheet=sheet, totals=totals)

    for row, amount in expected.items():
        previous_amount = 1000 if row == 8 else 0
        assert sheet.cell(row=row, column=4).value == previous_amount + amount

    assert sheet["D4"].value == 900
    assert sheet["B8"].value == 500
    assert sheet["B14"].value == "=B12+B13"

    workbook.close()
    print("A napi összesítések beírásának tesztje sikeres.")

def test_write_monthly_totals() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Október"

    for row in (82, 83, 84):
        sheet.merge_cells(
            start_row=row,
            end_row=row,
            start_column=2,
            end_column=3,
        )

    sheet["B82"] = 1000
    sheet["B81"] = "=SUM(B78:B80)"
    sheet["B74"] = 5000  # Kézzel kitöltött adat.

    totals = MonthlyTotals(
        year=2026,
        month=10,
        sany=500,
        diag=1000,
        kj=1500,
    )

    write_monthly_totals(sheet=sheet, totals=totals)

    assert sheet["B82"].value == 1500
    assert sheet["B83"].value == 1000
    assert sheet["B84"].value == 1500

    assert sheet["B81"].value == "=SUM(B78:B80)"
    assert sheet["B74"].value == 5000

    workbook.close()
    print("A havi összesítések beírásának tesztje sikeres.")

def test_write_import_result() -> None:
    template_workbook = Workbook()
    template = template_workbook.active
    template.title = "Szeptember"
    template["A82"] = "Segédanyag"

    workbook = create_year_workbook(
        template_sheet=template,
        year=2026,
        last_month=1,
    )

    workbook["Január"]["B82"] = 100
    workbook["Január"]["B74"] = 5000

    result = ImportResult(
        daily_totals=[
            DailyTotals(
                movement_date=date(2026, 2, 2),
                used_parts=14000,
                new_parts=37000,
                labor=22500,
                workshop_used_parts=10000,
                workshop_new_parts=33000,
                purchases=13125,
                shipping=1250,
                trailer=7500,
            ),
            DailyTotals(
                movement_date=date(2027, 3, 1),
                used_parts=999,
                new_parts=0,
                labor=0,
                workshop_used_parts=0,
                workshop_new_parts=0,
                purchases=0,
                shipping=0,
                trailer=0,
            ),
        ],
        monthly_totals=[
            MonthlyTotals(2026, 1, 500, 1000, 1500),
            MonthlyTotals(2027, 3, 999, 999, 999),
        ],
        new_documents={
            (2026, "TESZT-1"),
            (2027, "TESZT-2"),
        },
    )

    write_import_result(
        workbook=workbook,
        template_sheet=template,
        year=2026,
        result=result,
    )

    assert read_processed_documents(workbook) == {
        (2026, "TESZT-1")
    }

    assert "Február" in workbook.sheetnames
    assert "Március" not in workbook.sheetnames

    february = workbook["Február"]
    column = find_day_column(february, date(2026, 2, 2))
    assert february.cell(row=8, column=column).value == 14000
    assert february.cell(row=40, column=column).value == 7500

    assert workbook["Január"]["B82"].value == 600
    assert workbook["Január"]["B83"].value == 1000
    assert workbook["Január"]["B84"].value == 1500
    assert workbook["Január"]["B74"].value == 5000

    assert workbook["Összesítő"]["C2"].value == "Február"
    assert workbook["Euro"]["A3"].value == "Február"
    assert result.new_documents == {
        (2026, "TESZT-1"),
        (2027, "TESZT-2"),
    }

    workbook.close()
    template_workbook.close()
    print("Az importált napi és havi összesítések beírásának tesztje sikeres.")

def test_load_or_create_workbook() -> None:
    template_workbook = Workbook()
    template = template_workbook.active

    with TemporaryDirectory() as directory:
        path = get_year_workbook_path(directory, 2026)

        assert path.name == "Napi_értékesítési_riport_2026.xlsx"

        workbook = load_or_create_year_workbook(
            output_directory=directory,
            template_sheet=template,
            year=2026,
            last_month=2
        )

        assert "Január" in workbook.sheetnames
        assert "Február" in workbook.sheetnames
        assert "Március" not in workbook.sheetnames
        assert not path.exists() # A létrehozás még nem ment fájlt.

        workbook["Január"]["B8"] = 1234
        workbook["Január"]["B12"] = "=B8+B9+B10"
        workbook.save(path)
        workbook.close()

        loaded = load_or_create_year_workbook(
            output_directory=directory,
            template_sheet=template,
            year=2026,
        )

        assert loaded["Január"]["B8"].value == 1234
        assert loaded["Január"]["B12"].value == "=B8+B9+B10"
        assert "Február" in loaded.sheetnames

        loaded.close()

    template_workbook.close()
    print("A munkafüzet betöltésének vagy létrehozásának tesztje sikeres.")

def test_create_workbook_backup() -> None:
    with TemporaryDirectory() as directory:
        path = get_year_workbook_path(directory, 2026)

        # Nem létező riportról még nem készül másolat.
        assert create_workbook_backup(path) is None
        assert not (path.parent / "backups").exists()

        workbook = Workbook()
        workbook.active["B8"] = 1234
        workbook.active["B12"] = "=B8+B9+B10"
        workbook.save(path)
        workbook.close()

        original_content = path.read_bytes()

        backup_path = create_workbook_backup(path)

        assert backup_path is not None
        assert backup_path.exists()
        assert backup_path.parent.name == "backups"
        assert backup_path.name.startswith(f"{path.stem}_")
        assert backup_path.suffix == ".xlsx"

        # A másolat pontos, az eredeti fájl változatlan.
        assert backup_path.read_bytes() == original_content
        assert path.read_bytes() == original_content

        backup = load_workbook(backup_path, data_only=False)
        assert backup.active["B8"].value == 1234
        assert backup.active["B12"].value == "=B8+B9+B10"
        backup.close()
        print("A munkafüzet biztonsági mentésének tesztje sikeres.")

def test_save_workbook_safely() -> None:
    with TemporaryDirectory() as directory:
        path = get_year_workbook_path(directory, 2026)

        workbook = Workbook()
        sheet = workbook.active
        sheet["B8"] = 1000

        # Első mentés még nincs biztonsági másolat.
        assert save_workbook_safely(workbook, path) is None
        assert path.exists()

        sheet["B8"] = 2500
        backup_path = save_workbook_safely(workbook, path)
        
        assert backup_path is not None

        backup = load_workbook(backup_path)
        assert backup.active["B8"].value == 1000
        backup.close()

        saved = load_workbook(path)
        assert saved.active["B8"].value == 2500
        saved.close()

        original_content = path.read_bytes()

        # Félkész ideiglenes fájlt hagyó mentési hibát szimulálunk.
        def failing_save(filename) -> None:
            Path(filename).write_bytes(b"unfinished")
            raise OSError("Szimulált mentési hiba")

        with patch.object(workbook, "save", side_effect=failing_save):
            try:
                save_workbook_safely(workbook, path)
            except OSError as error:
                assert str(error) == "Szimulált mentési hiba"
            else:
                raise AssertionError("A mentési hiba nem terjed tovább.")

        assert path.read_bytes() == original_content
        assert not list(path.parent.glob(f".{path.stem}_*.xlsx"))

        # A biztonsági másolat hibája esetén sem írjuk felül a riportot.
        sheet["B8"] = 9000

        with patch(
            "axelpro_riport.excel_writer.create_workbook_backup",
            side_effect=OSError("Szimulált biztonsági másolat hiba")
        ):
            try:
                save_workbook_safely(workbook, path)
            except OSError as error:
                assert str(error) == "Szimulált biztonsági másolat hiba"
            else:
                raise AssertionError("A biztonsági másolat hibája nem terjed tovább.")

        assert path.read_bytes() == original_content
        assert not list(path.parent.glob(f".{path.stem}_*.xlsx"))
        
        workbook.close()
        print("A biztonságos mentés tesztje sikeres.")

def test_processed_documents() -> None:
    workbook = Workbook()

    assert read_processed_documents(workbook) == set()
    assert PROCESSED_DOCUMENTS_SHEET not in workbook.sheetnames

    sheet = ensure_processed_documents_sheet(workbook)

    assert sheet["A1"].value == "Év"
    assert sheet["B1"].value == "Bizonylatszám"
    assert sheet.sheet_state == "veryHidden"
    assert ensure_processed_documents_sheet(workbook) is sheet

    documents = {
        (2026, "SZ-2026/001"),
        (2026, "SZ-2026/002"),
        (2026, "=TESZT")
    }

    add_processed_documents(workbook, documents)

    assert read_processed_documents(workbook) == documents
    assert sheet.max_row == 4

    # Ugyanazok a kulcsok nem kerülnek be mégegyszer.
    add_processed_documents(workbook, documents)

    assert sheet.max_row == 4

    # A hibás kulcs nem módosíthatja a nyilvántartást.
    try:
        add_processed_documents(workbook, {(2026, "")})
    except ValueError:
        pass
    else:
        raise AssertionError("A hibás bizonylatkulcsot elfogadta.")

    assert read_processed_documents(workbook) == documents
    assert sheet.max_row == 4

    with TemporaryDirectory() as directory:
        path = get_year_workbook_path(directory, 2026)
        save_workbook_safely(workbook, path)

        loaded = load_workbook(path, data_only=False)
        try:
            assert read_processed_documents(loaded) == documents

            loaded_sheet = loaded[PROCESSED_DOCUMENTS_SHEET]
            assert loaded_sheet.sheet_state == "veryHidden"

            for row in loaded_sheet.iter_rows(min_row=2, max_col=2):
                assert row[1].data_type == "s"
        finally:
            loaded.close()

    workbook.close()
    print("A feldolgozott bizonylatok nyilvántartásának tesztje sikeres.")

def test_daily_exchange_rates() -> None:
    workbook = Workbook()

    day = date(2026, 10, 11)
    record = DailyExchangeRate(
        application_date=day,
        source_date=date(2026, 10, 9),
        rate=Decimal("395.12"),
    )
    rates = {day: record}

    try:
        assert read_daily_exchange_rates(workbook) == {}

        add_daily_exchange_rates(workbook, rates)

        sheet = workbook[DAILY_EXCHANGE_RATES_SHEET]

        assert sheet.sheet_state == "veryHidden"
        assert sheet["A2"].value == day
        assert sheet["B2"].value == date(2026, 10, 9)
        assert sheet["C2"].value == 395.12
        assert read_daily_exchange_rates(workbook) == rates

        # Az ismételt hozzáadás nem hoz létre új sort.
        add_daily_exchange_rates(workbook, rates)
        assert sheet.max_row == 2

        # Egy már alkalmazott árfolyam nem írható felül.
        conflicting_record = DailyExchangeRate(
            application_date=day,
            source_date=date(2026, 10, 9),
            rate=Decimal("396.00"),
        )

        try:
            add_daily_exchange_rates(
                workbook,
                {day: conflicting_record},
            )
        except ValueError:
            pass
        else:
            raise AssertionError("Az eltérő árfolyamot elfogadta.")

        assert read_daily_exchange_rates(workbook) == rates
        assert sheet.max_row == 2

        with TemporaryDirectory() as directory:
            path = get_year_workbook_path(directory, 2026)
            save_workbook_safely(workbook, path)

            loaded = load_workbook(path, data_only=False)
            try:
                assert read_daily_exchange_rates(loaded) == rates
                assert (
                    loaded[DAILY_EXCHANGE_RATES_SHEET].sheet_state
                    == "veryHidden"
                )
            finally:
                loaded.close()

    finally:
        workbook.close()

    print("A napi árfolyamok nyilvántartásának tesztje sikeres.")

def test_update_monthly_exchange_rates() -> None:
    workbook = Workbook()

    records = [
        DailyExchangeRate(
            application_date=date(2026, 10, 9),
            source_date=date(2026, 10, 9),
            rate=Decimal("395.12"),
        ),
        DailyExchangeRate(
            application_date=date(2026, 10, 30),
            source_date=date(2026, 10, 30),
            rate=Decimal("396.08"),
        ),
        # Novemberben alkalmazott, de októberi forrású árfolyam.
        DailyExchangeRate(
            application_date=date(2026, 11, 1),
            source_date=date(2026, 10, 30),
            rate=Decimal("396.08"),
        ),
        DailyExchangeRate(
            application_date=date(2026, 11, 2),
            source_date=date(2026, 11, 2),
            rate=Decimal("397.25"),
        ),
    ]

    try:
        add_daily_exchange_rates(
            workbook,
            {
                record.application_date: record
                for record in records
            },
        )

        # Egy forrásadat nélküli hónap értékét meg kell őrizni.
        euro = workbook.create_sheet("Euro")
        euro["A10"] = "Szeptember"
        euro["B10"] = 390
        euro["C10"] = date(2026, 9, 30)

        update_monthly_exchange_rates(workbook, year=2026)

        # Október sora: 10 + 1.
        assert euro["A11"].value == "Október"
        assert euro["B11"].value == 396.08
        assert euro["C11"].value == date(2026, 10, 30)

        # November sora: 11 + 1.
        assert euro["A12"].value == "November"
        assert euro["B12"].value == 397.25
        assert euro["C12"].value == date(2026, 11, 2)

        assert euro["B10"].value == 390
        assert euro["C10"].value == date(2026, 9, 30)

        # Az ismételt frissítés ugyanazt az eredményt adja.
        update_monthly_exchange_rates(workbook, year=2026)

        assert euro["B11"].value == 396.08
        assert euro["B12"].value == 397.25

    finally:
        workbook.close()

    print("A havi árfolyamok frissítésének tesztje sikeres.")

def run_all_tests() -> None:
    test_copy_block()
    test_clear_input_values()
    test_create_week_block()
    test_set_week_dates()
    test_create_month_sheet()
    test_ensure_month_sheet()
    test_summary_month()
    test_year_workbook()
    test_find_day_column()
    test_add_cell_amount()
    test_write_daily_totals()
    test_write_monthly_totals()
    test_write_import_result()
    test_load_or_create_workbook()
    test_create_workbook_backup()
    test_save_workbook_safely()
    test_processed_documents()
    test_daily_exchange_rates()
    test_update_monthly_exchange_rates()

if __name__ == "__main__":
    run_all_tests()
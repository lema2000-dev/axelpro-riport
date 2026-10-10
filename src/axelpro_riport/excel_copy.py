"""Cellák, formázások, méretek és összevonások másolása."""

from copy import copy

from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def copy_cell(
    source_sheet: Worksheet,
    target_sheet: Worksheet,
    source_row: int,
    source_column: int,
    target_row: int,
    target_column: int,
    *,
    translate_formulas: bool = True
) -> None:
    """Egy cella értékének, képletének és formázásának másolása."""
    source = source_sheet.cell(row=source_row, column=source_column)
    target = target_sheet.cell(row=target_row, column=target_column)

    if source.data_type == "f" and translate_formulas:
        target.value = Translator(
            source.value, origin=source.coordinate
        ).translate_formula(target.coordinate)
    else:
        target.value = source.value

    # A stílusobjektumokat másoljuk, így nem osztunk meg StyleProxy-t.
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def copy_range(
    source_sheet: Worksheet,
    target_sheet: Worksheet,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    target_start_row: int,
    target_start_column: int,
) -> None:
    """Egy téglalap alakú cellatartomány másolása."""
    for source_row in range(min_row, max_row + 1):
        for source_column in range(min_column, max_column + 1):
            target_row = target_start_row + (source_row - min_row)
            target_column = target_start_column + (source_column - min_column)

            copy_cell(
                source_sheet=source_sheet,
                target_sheet=target_sheet,
                source_row=source_row,
                source_column=source_column,
                target_row=target_row,
                target_column=target_column,
            )


def copy_row_heights(
    source_sheet: Worksheet,
    target_sheet: Worksheet,
    min_row: int,
    max_row: int,
    target_start_row: int,
) -> None:
    """A másolt sorok magasságának átvitele."""
    for source_row in range(min_row, max_row + 1):
        target_row = target_start_row + (source_row - min_row)
        target_sheet.row_dimensions[target_row].height = source_sheet.row_dimensions[
            source_row
        ].height


def copy_column_widths(
    source_sheet: Worksheet,
    target_sheet: Worksheet,
    min_column: int,
    max_column: int,
    target_start_column: int,
) -> None:
    """A másolt oszlopok szélességének átvitele."""
    for source_column in range(min_column, max_column + 1):
        target_column = target_start_column + (source_column - min_column)

        source_letter = get_column_letter(source_column)
        target_letter = get_column_letter(target_column)

        target_sheet.column_dimensions[target_letter].width = (
            source_sheet.column_dimensions[source_letter].width
        )


def copy_merged_ranges(
    source_sheet: Worksheet,
    target_sheet: Worksheet,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    target_start_row: int,
    target_start_column: int,
) -> None:
    """A másolt tartományon belüli cellaösszevonások átvitele."""
    row_offset = target_start_row - min_row
    column_offset = target_start_column - min_column

    for merged_range in list(source_sheet.merged_cells.ranges):
        if (
            merged_range.min_row >= min_row
            and merged_range.max_row <= max_row
            and merged_range.min_col >= min_column
            and merged_range.max_col <= max_column
        ):
            target_sheet.merge_cells(
                start_row=merged_range.min_row + row_offset,
                end_row=merged_range.max_row + row_offset,
                start_column=merged_range.min_col + column_offset,
                end_column=merged_range.max_col + column_offset,
            )


def copy_block(
    source_sheet: Worksheet,
    target_sheet: Worksheet,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    target_start_row: int,
    target_start_column: int,
) -> None:
    """Cellák, méretek és összevonások együttes másolása."""
    copy_range(
        source_sheet=source_sheet,
        target_sheet=target_sheet,
        min_row=min_row,
        max_row=max_row,
        min_column=min_column,
        max_column=max_column,
        target_start_row=target_start_row,
        target_start_column=target_start_column,
    )

    copy_row_heights(
        source_sheet=source_sheet,
        target_sheet=target_sheet,
        min_row=min_row,
        max_row=max_row,
        target_start_row=target_start_row,
    )

    copy_column_widths(
        source_sheet=source_sheet,
        target_sheet=target_sheet,
        min_column=min_column,
        max_column=max_column,
        target_start_column=target_start_column,
    )

    # Az összevonások csak a cellák másolása után kerüljenek át.
    copy_merged_ranges(
        source_sheet=source_sheet,
        target_sheet=target_sheet,
        min_row=min_row,
        max_row=max_row,
        min_column=min_column,
        max_column=max_column,
        target_start_row=target_start_row,
        target_start_column=target_start_column,
    )


def clear_input_values(
    sheet: Worksheet, min_row: int, max_row: int, min_column: int, max_column: int
) -> None:
    """A kézzel kitölthető cellák ürítése, a képletek megtartásával."""
    for row in sheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_column, max_col=max_column
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.data_type != "f":
                cell.value = None

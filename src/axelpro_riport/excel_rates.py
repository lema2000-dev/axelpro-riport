"""Napi árfolyamok Excel-tárolása és a havi Euro-tábla frissítése."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

from .excel_constants import DAILY_EXCHANGE_RATES_SHEET, MONTH_NAMES
from .models import DailyExchangeRate


def ensure_exchange_rate_sheet(workbook: WorkbookType) -> Worksheet:
    """Az Euro munkalap lekérése vagy létrehozása."""
    if "Euro" in workbook.sheetnames:
        return workbook["Euro"]

    sheet = workbook.create_sheet(title="Euro")

    sheet["A1"] = "Hónap"
    sheet["B1"] = "EUR/HUF árfolyam"

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 22

    return sheet


def ensure_month_exchange_rate(workbook: WorkbookType, month: int) -> str:
    """A havi árfolyam helyének előkészítése; cellahivatkozást ad vissza."""
    sheet = ensure_exchange_rate_sheet(workbook)
    row = month + 1

    sheet.cell(row=row, column=1).value = MONTH_NAMES[month - 1]
    # A korábban megadott árfolyam értékét nem írjuk felül.
    sheet.cell(row=row, column=2).number_format = "0.00"

    return f"'Euro'!$B${row}"


def ensure_daily_exchange_rates_sheet(workbook: WorkbookType) -> Worksheet:
    """A napi árfolyamokat tároló munkalap lekérése vagy létrehozása."""
    if DAILY_EXCHANGE_RATES_SHEET in workbook.sheetnames:
        sheet = workbook[DAILY_EXCHANGE_RATES_SHEET]
    else:
        sheet = workbook.create_sheet(title=DAILY_EXCHANGE_RATES_SHEET)

        sheet["A1"] = "Alkalmazás dátuma"
        sheet["B1"] = "Árfolyam dátuma"
        sheet["C1"] = "EUR/HUF árfolyam"

        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 22
        sheet.column_dimensions["C"].width = 22

    sheet.sheet_state = "veryHidden"

    return sheet


def read_daily_exchange_rates(
    workbook: WorkbookType,
) -> dict[date, DailyExchangeRate]:
    """A napi árfolyamok visszaolvasása a forrásdátumukkal együtt."""
    if DAILY_EXCHANGE_RATES_SHEET not in workbook.sheetnames:
        return {}

    sheet = workbook[DAILY_EXCHANGE_RATES_SHEET]

    expected_headers = (
        "Alkalmazás dátuma",
        "Árfolyam dátuma",
        "EUR/HUF árfolyam",
    )

    actual_headers = tuple(
        sheet.cell(row=1, column=column).value for column in range(1, 4)
    )

    if actual_headers != expected_headers:
        raise ValueError("A napi árfolyamok munkalapjának szerkezete nem megfelelő.")

    exchange_rates: dict[date, DailyExchangeRate] = {}

    for row_number, (application_date, source_date, value) in enumerate(
        sheet.iter_rows(
            min_row=2,
            max_col=3,
            values_only=True,
        ),
        start=2,
    ):
        if application_date is None and source_date is None and value is None:
            continue

        if isinstance(application_date, datetime):
            application_date = application_date.date()

        if isinstance(source_date, datetime):
            source_date = source_date.date()

        if (
            not isinstance(application_date, date)
            or not isinstance(source_date, date)
            or isinstance(value, bool)
            or not isinstance(value, (int, float, str, Decimal))
        ):
            raise ValueError(f"Hibás napi árfolyam a(z) {row_number}. sorban.")

        try:
            rate = Decimal(str(value))
        except InvalidOperation as exc:
            raise ValueError(f"Hibás napi árfolyam a(z) {row_number}. sorban.") from exc

        if application_date in exchange_rates:
            raise ValueError(
                f"Ismétlődő alkalmazási dátum: " f"{application_date:%Y-%m-%d}"
            )

        exchange_rates[application_date] = DailyExchangeRate(
            application_date=application_date,
            source_date=source_date,
            rate=rate,
        )

    return exchange_rates


def add_daily_exchange_rates(
    workbook: WorkbookType,
    exchange_rates: dict[date, DailyExchangeRate],
) -> None:
    """Új árfolyamrekordok hozzáadása a tárolt értékek megőrzésével."""
    existing_rates = read_daily_exchange_rates(workbook)
    new_rates: dict[date, DailyExchangeRate] = {}

    # Minden rekordot ellenőrzünk a munkalap módosítása előtt.
    for day, record in exchange_rates.items():
        if day != record.application_date:
            raise ValueError(f"Eltérő alkalmazási dátum: {day!r}")

        if day in existing_rates:
            if existing_rates[day] != record:
                raise ValueError(f"Eltérő tárolt árfolyamrekord: {day:%Y-%m-%d}")
            continue

        new_rates[day] = record

    if not new_rates:
        return

    sheet = ensure_daily_exchange_rates_sheet(workbook)

    for day, record in sorted(new_rates.items()):
        row = sheet.max_row + 1

        application_cell = sheet.cell(row=row, column=1)
        application_cell.value = day
        application_cell.number_format = "yyyy-mm-dd"

        source_cell = sheet.cell(row=row, column=2)
        source_cell.value = record.source_date
        source_cell.number_format = "yyyy-mm-dd"

        rate_cell = sheet.cell(row=row, column=3)
        rate_cell.value = float(record.rate)
        rate_cell.number_format = "0.00"


def update_monthly_exchange_rates(
    workbook: WorkbookType,
    year: int,
) -> None:
    """A havi árfolyamok frissítése a legutolsó tárolt forrásdátum alapján."""
    records = read_daily_exchange_rates(workbook)
    source_rates: dict[date, Decimal] = {}

    for record in records.values():
        source_date = record.source_date

        if source_date.year != year:
            continue

        if source_date in source_rates and source_rates[source_date] != record.rate:
            raise ValueError(
                f"Eltérő árfolyamok ugyanahhoz a forrásdátumhoz: "
                f"{source_date:%Y-%m-%d}"
            )

        source_rates[source_date] = record.rate

    if not source_rates:
        return

    sheet = ensure_exchange_rate_sheet(workbook)
    sheet["C1"] = "Árfolyam dátuma"
    sheet.column_dimensions["C"].width = 18

    months = sorted({day.month for day in source_rates})

    for month in months:
        latest_date = max(day for day in source_rates if day.month == month)

        ensure_month_exchange_rate(workbook, month)

        row = month + 1
        sheet.cell(row=row, column=2).value = float(source_rates[latest_date])

        date_cell = sheet.cell(row=row, column=3)
        date_cell.value = latest_date
        date_cell.number_format = "yyyy-mm-dd"


def get_monthly_exchange_rate_values(
    workbook: WorkbookType,
) -> tuple:
    """A havi árfolyamtábla értékei a változások összehasonlításához."""
    if "Euro" not in workbook.sheetnames:
        return ()

    return tuple(
        tuple(value.date() if isinstance(value, datetime) else value for value in row)
        for row in workbook["Euro"].iter_rows(
            min_row=1,
            max_row=13,
            max_col=3,
            values_only=True,
        )
    )

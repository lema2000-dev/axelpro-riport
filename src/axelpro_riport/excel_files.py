"""Sablonok és éves riportfájlok megnyitása, atomi mentése."""

from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

from .excel_layout import create_year_workbook


def load_template(template_path: str | Path) -> WorkbookType:
    """A végleges szeptemberi sablont tartalmazó munkafüzet betöltése."""
    path = Path(template_path)

    workbook = load_workbook(path, data_only=False)

    if "Szeptember" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("A sablon nem tartalmaz Szeptember nevű munkalapot.")

    return workbook


def get_year_workbook_path(output_directory: str | Path, year: int) -> Path:
    """Az adott év riportfájljának elérési útja."""
    if not 1 <= year <= 9999:
        raise ValueError("Az évszám 1 és 9999 közötti lehet.")

    return Path(output_directory) / f"Napi_értékesítési_riport_{year}.xlsx"


def load_or_create_year_workbook(
    output_directory: str | Path,
    template_sheet: Worksheet,
    year: int,
    *,
    last_month: int = 1,
) -> WorkbookType:
    """A meglévő éves riport megnyitása vagy új munkafüzet létrehozása."""
    path = get_year_workbook_path(output_directory, year)

    if path.exists():
        return load_workbook(path, data_only=False)

    return create_year_workbook(
        template_sheet=template_sheet, year=year, last_month=last_month
    )


def save_workbook_safely(workbook: WorkbookType, workbook_path: str | Path) -> None:
    """Mentés ideiglenes fájlba, majd a kész riport atomi cseréje."""
    path = Path(workbook_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    with NamedTemporaryFile(
        dir=path.parent, prefix=f".{path.stem}_", suffix=".xlsx", delete=False
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        workbook.save(temporary_path)

        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)

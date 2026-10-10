"""Az importált bizonylatok rejtett Excel-nyilvántartása."""

from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

from .excel_constants import PROCESSED_DOCUMENTS_SHEET


def ensure_processed_documents_sheet(workbook: WorkbookType) -> Worksheet:
    """A bizonylatnyilvántartó munkalap lekérése vagy létrehozása."""
    if PROCESSED_DOCUMENTS_SHEET in workbook.sheetnames:
        sheet = workbook[PROCESSED_DOCUMENTS_SHEET]
    else:
        sheet = workbook.create_sheet(title=PROCESSED_DOCUMENTS_SHEET)

        sheet["A1"] = "Év"
        sheet["B1"] = "Bizonylatszám"

    sheet.sheet_state = "veryHidden"

    return sheet


def read_processed_documents(workbook: WorkbookType) -> set[tuple[int, str]]:
    """A nyilvántartott bizonylatkulcsok visszaolvasása."""
    if PROCESSED_DOCUMENTS_SHEET not in workbook.sheetnames:
        return set()

    sheet = workbook[PROCESSED_DOCUMENTS_SHEET]
    documents: set[tuple[int, str]] = set()

    for row_number, (year, document_number) in enumerate(
        sheet.iter_rows(min_row=2, max_col=2, values_only=True),
        start=2,
    ):
        if year is None and document_number is None:
            continue

        if (
            isinstance(year, bool)
            or not isinstance(year, int)
            or not 1 <= year <= 9999
            or not isinstance(document_number, str)
            or not document_number.strip()
        ):
            raise ValueError(
                f"Hibás bizonylatnyilvántartás a(z) {sheet.title} munkalap {row_number}. sorában."
            )

        documents.add((year, document_number))

    return documents


def add_processed_documents(
    workbook: WorkbookType, documents: set[tuple[int, str]]
) -> None:
    """Az új bizonylatkulcsok hozzáadása a nyilvántartáshoz."""
    existing_documents = read_processed_documents(workbook)
    new_documents = documents - existing_documents

    if not new_documents:
        return

    sheet = ensure_processed_documents_sheet(workbook)

    for year, document_number in sorted(new_documents):
        row = sheet.max_row + 1

        sheet.cell(row=row, column=1).value = year

        cell = sheet.cell(row=row, column=2)
        cell.value = document_number
        cell.data_type = "s"
